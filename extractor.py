
import csv
import os
import docx
import openpyxl
import pytesseract
from PIL import Image
from pypdf import PdfReader
import config

if os.path.exists(config.TESSERACT_PATH):
    pytesseract.pytesseract.tesseract_cmd = config.TESSERACT_PATH


def extract_from_pdf(path):
    reader = PdfReader(path)
    parts = []
    for page in reader.pages:
        text = page.extract_text()
        if text:
            parts.append(text)
    return "\n".join(parts)


def extract_from_image(path):
    if not os.path.exists(config.TESSERACT_PATH):
        raise RuntimeError("Tesseract is not installed, cannot read images")
    image = Image.open(path)
    return pytesseract.image_to_string(image)


def extract_from_text(path):
    with open(path, "r", encoding="utf-8", errors="ignore") as f:
        return f.read()


def extract_from_docx(path):
    document = docx.Document(path)
    parts = []
    for para in document.paragraphs:
        if para.text.strip():
            parts.append(para.text)
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            line = " | ".join(cells)
            if line.strip(" |"):
                parts.append(line)
    return "\n".join(parts)


def extract_from_excel(path):
    book = openpyxl.load_workbook(path, data_only=True)
    parts = []

    for sheet in book.worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            continue

        parts.append("Sheet: " + sheet.title)

        header = []
        for value in rows[0]:
            header.append("" if value is None else str(value).strip())

        parts.append(" | ".join(header))

        for row in rows[1:]:
            cells = []
            for value in row:
                cells.append("" if value is None else str(value).strip())
            if not any(cells):
                continue
            # keep the column name next to each value so a row reads on its own
            labelled = []
            for i, value in enumerate(cells):
                if not value:
                    continue
                name = header[i] if i < len(header) and header[i] else "Column " + str(i + 1)
                labelled.append(name + ": " + value)
            if labelled:
                parts.append(", ".join(labelled))

        parts.append("")

    book.close()
    return "\n".join(parts)


def extract_from_csv(path):
    parts = []
    with open(path, "r", encoding="utf-8", errors="ignore", newline="") as f:
        reader = csv.reader(f)
        rows = list(reader)

    if not rows:
        return ""

    header = [c.strip() for c in rows[0]]
    parts.append(" | ".join(header))

    for row in rows[1:]:
        labelled = []
        for i, value in enumerate(row):
            value = value.strip()
            if not value:
                continue
            name = header[i] if i < len(header) and header[i] else "Column " + str(i + 1)
            labelled.append(name + ": " + value)
        if labelled:
            parts.append(", ".join(labelled))

    return "\n".join(parts)


def extract_text(path):
    ext = os.path.splitext(path)[1].lower()

    if ext == ".pdf":
        text = extract_from_pdf(path)
    elif ext in (".png", ".jpg", ".jpeg"):
        text = extract_from_image(path)
    elif ext in (".txt", ".md"):
        text = extract_from_text(path)
    elif ext == ".docx":
        text = extract_from_docx(path)
    elif ext in (".xlsx", ".xls"):
        text = extract_from_excel(path)
    elif ext == ".csv":
        text = extract_from_csv(path)
    else:
        raise ValueError("Unsupported file type: " + ext)

    return clean_text(text)


def clean_text(text):
    lines = []
    for line in text.splitlines():
        line = line.strip()
        if line:
            lines.append(line)
    return "\n".join(lines)